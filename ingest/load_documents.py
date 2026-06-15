import os
from typing import List

from langchain_core.documents import Document
from langchain_community.document_loaders import PyPDFLoader, TextLoader


def load_pdf(path: str) -> List[Document]:
    if not os.path.exists(path):
        raise FileNotFoundError(path)
    return PyPDFLoader(path).load()


def load_txt(path: str) -> List[Document]:
    if not os.path.exists(path):
        raise FileNotFoundError(path)
    return TextLoader(path, encoding="utf-8").load()


def load_docx(path: str) -> List[Document]:
    if not os.path.exists(path):
        raise FileNotFoundError(path)
    try:
        from docx import Document as DocxDocument

        doc = DocxDocument(path)
        text = "\n".join(para.text for para in doc.paragraphs if para.text.strip())
        return [Document(page_content=text, metadata={"source": path})]
    except ImportError as exc:
        raise ImportError("Install python-docx to load .docx files: pip install python-docx") from exc


def load_documents_from_paths(paths: list[str]) -> List[Document]:
    docs: List[Document] = []
    for path in paths:
        ext = os.path.splitext(path)[1].lower()
        if ext == ".pdf":
            docs.extend(load_pdf(path))
        elif ext in (".txt", ".md"):
            docs.extend(load_txt(path))
        elif ext == ".docx":
            docs.extend(load_docx(path))
        else:
            raise ValueError(f"Unsupported file type: {path}")
    return docs
